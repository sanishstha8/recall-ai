from __future__ import annotations

from typing import IO

import docx
from docx.opc.exceptions import PackageNotFoundError

from apps.documents.parsers.base import (
    BaseDocumentParser,
    DocumentParsingError,
    ParsedDocument,
    pages_from_texts,
)
from apps.documents.processing.cleaning import paginate_text


class DOCXParser(BaseDocumentParser):
    file_types = ("docx",)

    def parse(self, file: IO[bytes]) -> ParsedDocument:
        try:
            document = docx.Document(file)
        except (PackageNotFoundError, KeyError, ValueError) as exc:
            raise DocumentParsingError(f"Could not read the Word document: {exc}") from exc

        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            if style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                text = f"{'#' * min(int(level), 6)} {text}"
            elif "list" in style:
                text = f"- {text}"
            blocks.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        return ParsedDocument(pages=pages_from_texts(paginate_text("\n\n".join(blocks))), native_pages=False)
