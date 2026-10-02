import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.documents.models import Document, DocumentChunk, DocumentStatus, ProcessingStage
from apps.documents.processing.chunking import chunk_pages
from apps.documents.processing.cleaning import clean_text
from apps.knowledge.models import KnowledgeBase
from apps.usage.models import Operation, UsageRecord
from apps.workspaces.models import Role, WorkspaceMember
from common.storage import get_document_storage
from common.utils import sanitize_filename
from tests.conftest import client_for
from tests.pdf_factory import build_pdf

pytestmark = pytest.mark.django_db

DB_PAGES = [
    "Chapter 4: Normalization. Normalization organizes relational data to reduce redundancy and "
    "update anomalies. A relation is in first normal form (1NF) when every attribute holds atomic values.",
    "Second normal form (2NF) requires 1NF and that no non-key attribute depends on part of a composite key. "
    "Third normal form (3NF) removes transitive dependencies between non-key attributes.",
    "Chapter 5: Transactions. ACID stands for atomicity, consistency, isolation and durability.",
]


@pytest.fixture
def kb(workspace):
    return KnowledgeBase.objects.create(workspace=workspace, name="BSc CSIT")


@pytest.fixture
def upload(client, django_capture_on_commit_callbacks):
    def _upload(kb, filename="Advanced Database.pdf", content=None, c=None):
        content = build_pdf(DB_PAGES) if content is None else content
        with django_capture_on_commit_callbacks(execute=True):
            return (c or client).post(
                f"/api/v1/knowledge-bases/{kb.id}/documents/",
                {"file": SimpleUploadedFile(filename, content)},
                format="multipart",
            )

    return _upload


def test_pdf_upload_is_processed_end_to_end(client, kb, upload):
    res = upload(kb)
    assert res.status_code == 201, res.data
    assert res.data["name"] == "Advanced Database.pdf"

    document = Document.objects.get(pk=res.data["id"])
    assert document.status == DocumentStatus.READY, document.processing_error
    assert document.processing_stage == ProcessingStage.DONE
    assert document.page_count == 3
    assert document.word_count > 40
    assert document.processed_at is not None

    pages = list(document.pages.order_by("page_number"))
    assert [p.page_number for p in pages] == [1, 2, 3]
    assert "first normal form" in pages[0].content

    chunks = list(document.chunks.select_related("page"))
    assert len(chunks) == document.chunk_count >= 3
    for chunk in chunks:
        assert chunk.embedding is not None and len(chunk.embedding) == chunk.embedding_dimensions
        assert chunk.knowledge_base_id == kb.id
        assert chunk.metadata["page_number"] == chunk.page.page_number
        assert chunk.metadata["document_id"] == document.id
        assert chunk.content in chunk.page.content or chunk.page.content in chunk.content

    usage = UsageRecord.objects.get(operation=Operation.EMBEDDING)
    assert usage.workspace_id == kb.workspace_id and usage.total_tokens > 0

    status_res = client.get(f"/api/v1/documents/{document.id}/status/")
    assert status_res.data["status"] == "READY"
    page_res = client.get(f"/api/v1/documents/{document.id}/pages/", {"page_number": 2})
    assert "Third normal form" in page_res.data["content"]
    search = client.get(f"/api/v1/documents/{document.id}/search-text/", {"q": "normal form"})
    assert [r["page_number"] for r in search.data] == [1, 2]


def test_storage_path_never_uses_uploaded_filename(kb, upload):
    res = upload(kb, filename="../../etc/passwd.pdf")
    document = Document.objects.get(pk=res.data["id"])
    assert ".." not in document.storage_path and "passwd" not in document.storage_path
    assert document.original_filename == "passwd.pdf"
    assert get_document_storage().exists(document.storage_path)


def test_file_download_is_scoped(client, other_client, kb, upload):
    document_id = upload(kb).data["id"]
    res = client.get(f"/api/v1/documents/{document_id}/file/")
    assert res.status_code == 200
    assert res["Content-Type"] == "application/pdf"
    assert res["X-Content-Type-Options"] == "nosniff"
    assert b"".join(res.streaming_content).startswith(b"%PDF")
    assert other_client.get(f"/api/v1/documents/{document_id}/file/").status_code == 404


@pytest.mark.parametrize(
    ("filename", "content", "message"),
    [
        ("notes.exe", b"MZ\x90\x00", "Unsupported file type"),
        ("fake.pdf", b"<html>not a pdf</html>", "does not match"),
        ("binary.txt", b"\x00\x01\x02\x03", "does not match"),
        ("fake.docx", b"PK\x03\x04garbage", "does not match"),
        ("empty.pdf", b"", "empty"),
    ],
)
def test_upload_validation(kb, upload, filename, content, message):
    res = upload(kb, filename=filename, content=content)
    assert res.status_code == 400
    assert message in res.data["error"]["message"]
    assert not Document.objects.exists()


def test_upload_size_limit(kb, upload, settings):
    settings.MAX_UPLOAD_SIZE = 100
    res = upload(kb)
    assert res.status_code == 400
    assert "MB or smaller" in res.data["error"]["message"]


def test_text_markdown_and_docx_parsing(kb, upload):
    import docx

    md = upload(kb, "notes.md", b"# Indexes\n\nB-tree indexes speed up range queries.\n\n- hash index\n")
    txt = upload(kb, "notes.txt", "Plain text about SQL joins — inner and outer.".encode())
    buffer = io.BytesIO()
    d = docx.Document()
    d.add_heading("Concurrency Control", level=1)
    d.add_paragraph("Two-phase locking guarantees serializability.")
    d.save(buffer)
    wd = upload(kb, "lecture.docx", buffer.getvalue())

    for res, needle in [(md, "# Indexes"), (txt, "inner and outer"), (wd, "# Concurrency Control")]:
        document = Document.objects.get(pk=res.data["id"])
        assert document.status == DocumentStatus.READY, document.processing_error
        assert needle in document.pages.first().content


def test_pdf_without_text_fails_with_reason(kb, upload, client):
    res = upload(kb, content=build_pdf(["", ""]))
    document = Document.objects.get(pk=res.data["id"])
    assert document.status == DocumentStatus.FAILED
    assert "No extractable text" in document.processing_error
    detail = client.get(f"/api/v1/documents/{document.id}/")
    assert detail.data["processing_error"] == document.processing_error


def test_embedding_outage_marks_document_failed_and_reprocess_recovers(
    kb, upload, client, settings, django_capture_on_commit_callbacks
):
    settings.AI_EMBEDDING_PROVIDER = "tests.test_documents.DownEmbeddingProvider"
    document_id = upload(kb).data["id"]
    document = Document.objects.get(pk=document_id)
    assert document.status == DocumentStatus.FAILED
    assert "Cannot reach Ollama" in document.processing_error
    assert not document.chunks.exists()

    settings.AI_EMBEDDING_PROVIDER = "tests.fakes.HashingEmbeddingProvider"
    with django_capture_on_commit_callbacks(execute=True):
        res = client.post(f"/api/v1/documents/{document_id}/reprocess/")
    assert res.status_code == 202
    document.refresh_from_db()
    assert document.status == DocumentStatus.READY
    assert document.processing_error == ""


def test_reprocess_rebuilds_chunks_without_duplicates(kb, upload, client, django_capture_on_commit_callbacks):
    document_id = upload(kb).data["id"]
    before = DocumentChunk.objects.filter(document_id=document_id).count()
    with django_capture_on_commit_callbacks(execute=True):
        client.post(f"/api/v1/documents/{document_id}/reprocess/")
    assert DocumentChunk.objects.filter(document_id=document_id).count() == before


def test_reprocess_conflict_while_processing(kb, upload, client):
    document = Document.objects.get(pk=upload(kb).data["id"])
    Document.objects.filter(pk=document.pk).update(status=DocumentStatus.PROCESSING)
    res = client.post(f"/api/v1/documents/{document.pk}/reprocess/")
    assert res.status_code == 409


def test_delete_removes_file_and_chunks(kb, upload, client, django_capture_on_commit_callbacks):
    document = Document.objects.get(pk=upload(kb).data["id"])
    with django_capture_on_commit_callbacks(execute=True):
        assert client.delete(f"/api/v1/documents/{document.id}/").status_code == 204
    assert not DocumentChunk.objects.filter(document_id=document.id).exists()
    assert not get_document_storage().exists(document.storage_path)


def test_document_isolation(kb, upload, other_client, other_user):
    document_id = upload(kb).data["id"]
    assert other_client.get(f"/api/v1/documents/{document_id}/").status_code == 404
    assert other_client.get(f"/api/v1/documents/{document_id}/pages/").status_code == 404
    assert other_client.delete(f"/api/v1/documents/{document_id}/").status_code == 404
    assert other_client.post(f"/api/v1/documents/{document_id}/reprocess/").status_code == 404
    assert other_client.get(f"/api/v1/knowledge-bases/{kb.id}/documents/").status_code == 404
    assert other_client.get("/api/v1/documents/").data["results"] == []
    res = other_client.post(
        f"/api/v1/knowledge-bases/{kb.id}/documents/",
        {"file": SimpleUploadedFile("x.pdf", build_pdf(["hi"]))},
        format="multipart",
    )
    assert res.status_code == 404


def test_viewer_cannot_upload_or_delete(kb, upload, workspace, other_user):
    document_id = upload(kb).data["id"]
    WorkspaceMember.objects.create(workspace=workspace, user=other_user, role=Role.VIEWER)
    viewer = client_for(other_user)
    assert viewer.get(f"/api/v1/documents/{document_id}/").status_code == 200
    assert upload(kb, c=viewer).status_code == 403
    assert viewer.delete(f"/api/v1/documents/{document_id}/").status_code == 403


def test_clean_text():
    raw = "Normal-\nization   is\x00 key.\r\n\r\n\r\n\r\nNext" + chr(0xA0) + "para"
    assert clean_text(raw) == "Normalization is key.\n\nNext para"


def test_chunking_respects_size_overlap_and_pages(settings):
    page_one = "\n\n".join(f"Sentence number {i} about relational algebra and joins." for i in range(60))
    chunks = chunk_pages([(1, page_one), (2, "Short second page.")], chunk_size=120, overlap=30)
    assert {c.page_number for c in chunks} == {1, 2}
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    assert all(c.token_count <= 120 for c in chunks)
    page_one_chunks = [c for c in chunks if c.page_number == 1]
    assert len(page_one_chunks) > 3
    # Consecutive chunks overlap: the next chunk starts with the previous chunk's trailing units.
    first_units = page_one_chunks[0].content.split("\n\n")
    second_head = page_one_chunks[1].content.split("\n\n")[0]
    assert second_head in first_units[-3:]
    assert first_units[-1] in page_one_chunks[1].content
    assert chunks[-1].content == "Short second page."


def test_chunking_splits_giant_paragraph():
    text = " ".join(["word"] * 3000)
    chunks = chunk_pages([(1, text)], chunk_size=200, overlap=20)
    assert len(chunks) > 10
    assert all(c.token_count <= 200 for c in chunks)


def test_sanitize_filename():
    assert sanitize_filename("../../a/b\\c<script>.pdf") == "c_script_.pdf"
    assert sanitize_filename("   ") == "document"


class DownEmbeddingProvider:
    model = "down"
    dimensions = 768

    def embed_documents(self, texts):
        from apps.ai.providers import AIProviderUnavailable

        raise AIProviderUnavailable("Cannot reach Ollama at http://ollama:11434.")
